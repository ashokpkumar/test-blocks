"""Turns a block program into real Selenium actions.

Every block type is one function registered with @block("type"). It receives
the Run (browser + helpers), the block's fields and its inner blocks (for
"repeat" / "if"), and returns a short, kid-friendly sentence about what it did.

Scrape blocks remember what they grab under a name; any text field can use it
again by writing {name}.
"""
import base64
import csv
import io
import ipaddress
import json
import os
import re
import socket
import random
import threading
import time
import urllib.request
from datetime import datetime
from pathlib import Path
from urllib.parse import urlparse

from selenium import webdriver
from selenium.common.exceptions import (
    ElementClickInterceptedException,
    ElementNotInteractableException,
    InvalidSelectorException,
    NoSuchElementException,
    StaleElementReferenceException,
    TimeoutException,
    WebDriverException,
)
from selenium.webdriver.common.action_chains import ActionChains
from selenium.webdriver.common.by import By
from selenium.webdriver.common.keys import Keys
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.support.ui import Select, WebDriverWait

LOCATORS = {
    "css": By.CSS_SELECTOR,
    "xpath": By.XPATH,
    "id": By.ID,
    "name": By.NAME,
    "link": By.LINK_TEXT,
    "partial": By.PARTIAL_LINK_TEXT,
    "class": By.CLASS_NAME,
    "tag": By.TAG_NAME,
}

KEYS = {
    "enter": Keys.ENTER,
    "tab": Keys.TAB,
    "escape": Keys.ESCAPE,
    "space": Keys.SPACE,
    "backspace": Keys.BACKSPACE,
    "delete": Keys.DELETE,
    "up": Keys.ARROW_UP,
    "down": Keys.ARROW_DOWN,
    "left": Keys.ARROW_LEFT,
    "right": Keys.ARROW_RIGHT,
    "home": Keys.HOME,
    "end": Keys.END,
    "pageup": Keys.PAGE_UP,
    "pagedown": Keys.PAGE_DOWN,
}

BLOCKS = {}


def block(*types):
    def register(fn):
        for t in types:
            BLOCKS[t] = fn
        return fn
    return register


class Oops(Exception):
    """A step failed in a way we can explain in plain words."""


class Stopped(Exception):
    """The kid pressed Stop."""


class StepFailed(Exception):
    """A step already reported its failure; outer blocks just mark themselves."""


def friendly(e):
    if isinstance(e, Oops):
        return str(e)
    if isinstance(e, TimeoutException):
        return "That took too long, so I gave up."
    if isinstance(e, ElementClickInterceptedException):
        return "Something else is covering that element, so I can't click it."
    if isinstance(e, ElementNotInteractableException):
        return "I found it, but I can't use it (it may be hidden or disabled)."
    if isinstance(e, InvalidSelectorException):
        return "That element address doesn't look right. Try #id, .class or a tag name with css."
    if isinstance(e, StaleElementReferenceException):
        return "The page changed while I was using that element. Try adding a wait."
    if isinstance(e, WebDriverException):
        return (e.msg or str(e)).strip().splitlines()[0]
    return f"{type(e).__name__}: {e}"


def web_address(url):
    """Kids type "example.com" or "localhost:8000"; add the missing http(s)://."""
    url = str(url).strip()
    if "://" in url or url.startswith(("about:", "data:")):
        return url
    local = url.startswith(("localhost", "127.0.0.1", "0.0.0.0", "192.168.", "10."))
    return ("http://" if local else "https://") + url


def num(value, what="number"):
    try:
        return float(str(value).strip())
    except ValueError:
        raise Oops(f'"{value}" is not a {what}. Please type digits.')


def as_text(value):
    """How a remembered value looks when it is put into a text field with {name}."""
    if isinstance(value, dict):
        return as_text(value["rows"])
    if isinstance(value, list):
        return ", ".join(" | ".join(map(str, v)) if isinstance(v, list) else str(v) for v in value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def var_name(name):
    return re.sub(r"\W+", "_", str(name).strip().strip("{}")).strip("_")


def safe_filename(name):
    stem, dot, ext = str(name).rpartition(".")
    clean = lambda s: "".join(c for c in s if c.isalnum() or c in "-_ ").strip() or "file"
    return f"{clean(stem)}.{clean(ext)}" if dot else clean(name)


def as_number(value):
    """Numbers stay numbers; whole floats become ints so {count} reads 3, not 3.0."""
    n = value if isinstance(value, (int, float)) else number_in(as_text(value))
    return int(n) if float(n).is_integer() else n


def rows_of(name, value):
    """Any variable as spreadsheet rows (header first)."""
    if isinstance(value, dict):
        return [value["columns"], *value["rows"]]
    if isinstance(value, list):
        return [[name], *([v] if not isinstance(v, list) else v for v in value)]
    return [[name], [value]]


def to_csv(rows):
    out = io.StringIO()
    csv.writer(out).writerows(rows)
    return "﻿" + out.getvalue()  # BOM so Excel shows emoji/accents


def to_xlsx(sheets):
    """sheets: {sheet name: rows}. First row of each sheet is bold."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    wb.remove(wb.active)
    for title, rows in sheets.items():
        ws = wb.create_sheet(re.sub(r"[\[\]:*?/\\]", "_", title)[:31] or "Sheet")
        for row in rows:  # number-looking text becomes a real number, so Excel can add it up
            ws.append([v if isinstance(v, (int, float)) else as_number(v) if re.fullmatch(r"-?\d+(\.\d+)?", str(v).strip()) else str(v) for v in row])
        for cell in ws[1]:
            cell.font = Font(bold=True)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def read_table(name, data):
    """A loaded file as a variable: CSV/Excel -> table, JSON -> as is, text -> list of lines."""
    ext = name.rsplit(".", 1)[-1].lower()
    if ext in ("xlsx", "xlsm"):
        from openpyxl import load_workbook
        rows = [["" if c is None else c for c in r] for r in load_workbook(io.BytesIO(data), read_only=True, data_only=True).active.iter_rows(values_only=True)]
    elif ext in ("csv", "tsv"):
        rows = list(csv.reader(io.StringIO(data.decode("utf-8-sig")), delimiter="\t" if ext == "tsv" else ","))
    elif ext == "json":
        return json.loads(data.decode("utf-8-sig"))
    else:
        return [line for line in data.decode("utf-8-sig", errors="replace").splitlines() if line.strip()]
    rows = [r for r in rows if any(str(c).strip() for c in r)]
    if not rows:
        raise Oops(f"{name} is empty.")
    return {"columns": [str(c) for c in rows[0]], "rows": [[as_text(c) for c in r] for r in rows[1:]]}


def number_in(text):
    """Pull the number out of things like "$12.99" or "3 items"."""
    found = re.search(r"-?\d[\d,]*(\.\d+)?", str(text))
    if not found:
        raise Oops(f'I couldn\'t find a number in "{text}".')
    return float(found[0].replace(",", ""))


class Run:
    """One test run.

    server_mode is for the copy hosted on becomeaiengg.com: strangers drive the
    browser, so it is always headless Chromium, private addresses and server
    files are off limits, runs are time-limited, and screenshots / saved files
    go back to the visitor instead of onto the server's disk.
    """

    def __init__(self, emit, step_delay=0.4, keep_open=False, timeout=10,
                 shots_dir=Path("screenshots"), scraped_dir=Path("scraped"),
                 server_mode=False, port=8000, max_seconds=None):
        self.emit = emit
        self.step_delay = step_delay
        self.server_mode = server_mode
        self.keep_open = keep_open and not server_mode
        self.timeout = timeout
        self.shots_dir = Path(shots_dir)
        self.scraped_dir = Path(scraped_dir)
        self.port = port
        self.deadline = time.time() + max_seconds if max_seconds else None
        self.max_seconds = max_seconds
        self.out_dir = self.scraped_dir
        self.driver = None
        self.live = False
        self._last_frame = 0
        self._answer = None
        self._answered = threading.Event()
        self.passed = 0
        self.vars = {}
        self._stop = threading.Event()

    # ---------- life cycle ----------

    def stop(self):
        self._stop.set()

    def check_stop(self):
        if self._stop.is_set():
            raise Stopped()
        if self.deadline and time.time() > self.deadline:
            limit = f"{self.max_seconds // 60} minutes" if self.max_seconds >= 120 else f"{self.max_seconds} seconds"
            raise Oops(f"The test ran longer than {limit}, so I stopped it.")

    def pause(self, seconds):
        end = time.time() + seconds
        while time.time() < end:
            self.check_stop()
            self.frame(force=False)
            time.sleep(min(0.1, max(0, end - time.time())))

    def log(self, text):
        self.emit({"event": "log", "text": text})

    def execute(self, scripts):
        ok, stopped, text = False, False, ""
        try:
            for blocks in scripts:
                self.run_blocks(blocks)
            ok = True
            text = f"All {self.passed} steps passed!"
        except StepFailed as e:
            text = str(e)
        except Stopped:
            stopped, text = True, "Stopped."
        except Exception as e:  # anything unexpected still ends the run cleanly
            text = friendly(e)
        finally:
            if self.driver and not self.keep_open:
                try:
                    self.driver.quit()
                except Exception:
                    pass
            self.emit({"event": "done", "ok": ok, "stopped": stopped, "text": text, "passed": self.passed})

    def run_blocks(self, blocks):
        for b in blocks:
            self.check_stop()
            bid, kind = b.get("id"), b.get("type")
            self.emit({"event": "step", "id": bid, "status": "running"})
            try:
                fn = BLOCKS.get(kind)
                if fn is None:
                    raise Oops(f'I don\'t know how to do the "{kind}" block yet.')
                self.current_raw = b.get("fields") or {}  # unfilled, for loops that re-check {variables}
                fields = {k: self.fill(v) for k, v in self.current_raw.items()}
                message = fn(self, fields, b.get("children") or [])
            except Stopped:
                self.emit({"event": "step", "id": bid, "status": "stopped"})
                raise
            except StepFailed:
                self.emit({"event": "step", "id": bid, "status": "fail"})
                raise
            except Exception as e:
                text = friendly(e)
                self.emit({"event": "step", "id": bid, "status": "fail", "text": text})
                self.frame()
                raise StepFailed(text)
            if kind != "start":
                self.passed += 1
            self.emit({"event": "step", "id": bid, "status": "pass", "text": message})
            self.frame()
            self.pause(self.step_delay)

    def frame(self, force=True):
        """Send a picture of the hidden browser (plus its address) for the little
        browser window in the editor. Called after every step and, throttled,
        while waiting, so kids can watch pages load and things appear."""
        if not (self.live and self.driver):
            return
        now = time.time()
        if not force and now - self._last_frame < 0.5:
            return
        self._last_frame = now
        try:
            # Ask for the address first: while an alert() is open this raises
            # (prompts are set to "ignore", so the popup stays open for the next
            # block) and we skip the frame instead of hanging on the screenshot.
            url, title = self.driver.current_url, self.driver.title
            shot = self.driver.execute_cdp_cmd("Page.captureScreenshot", {"format": "jpeg", "quality": 55})
            self.emit({"event": "frame", "src": "data:image/jpeg;base64," + shot["data"], "url": url, "title": title})
        except Exception:
            pass  # e.g. a popup is open; the next step sends a new frame

    def check_address(self, url):
        """On the website, keep the browser on the public internet (and our own playground)."""
        if not self.server_mode:
            return
        p = urlparse(url)
        if p.scheme not in ("http", "https"):
            raise Oops("On the website I can only visit http:// and https:// pages.")
        host = p.hostname or ""
        if host in ("localhost", "127.0.0.1") and (p.port or 80) == self.port and p.path.startswith("/playground"):
            return
        try:
            addresses = {info[4][0] for info in socket.getaddrinfo(host, p.port or 443)}
        except socket.gaierror:
            raise Oops(f'I couldn\'t find a website called "{host}". Check the address.')
        for a in addresses:
            ip = ipaddress.ip_address(a.split("%")[0])
            if not ip.is_global:
                raise Oops(f'"{host}" is a private address, so I\'m not allowed to visit it from the website.')

    def send_file(self, name, data, mime):
        """Website: hand the file to the visitor's browser. Locally: it's already on disk."""
        payload = data if isinstance(data, bytes) else data.encode("utf-8")
        self.emit({"event": "download", "name": name, "mime": mime,
                   "data": base64.b64encode(payload).decode("ascii")})

    # ---------- remembered (scraped) values ----------

    def fill(self, value):
        """Swap {name} for a remembered value. Unknown names are left alone."""
        if not isinstance(value, str) or "{" not in value:
            return value
        return re.sub(r"\{(\w+)\}", lambda m: as_text(self.vars[m[1]]) if m[1] in self.vars else m[0], value)

    def remember(self, name, value, card=False):
        """Store a variable. Every change shows up in the editor's 📦 Variables box;
        card=True (scrape blocks) also puts a preview card in the diary."""
        name = var_name(name)
        if not name:
            raise Oops("Pick a variable for this block, like price or titles.")
        self.vars[name] = value
        kind = "table" if isinstance(value, dict) else "list" if isinstance(value, list) else "text"
        self.emit({"event": "var", "name": name, "kind": kind, "value": value})
        if card:
            self.emit({"event": "scraped", "name": name, "kind": kind, "value": value})
        return name

    def recall(self, name):
        name = var_name(name)
        if name not in self.vars:
            known = ", ".join(self.vars) or "nothing yet"
            raise Oops(f'The variable "{name}" has nothing in it yet (variables with values: {known}).')
        return name, self.vars[name]

    def output(self, filename, data, mime):
        """Hand a finished file to the kid: a download on the website, a file in
        the output folder on your own computer."""
        filename = safe_filename(filename)
        if self.server_mode:
            self.send_file(filename, data, mime)
            return f"Downloaded {filename}"
        self.out_dir.mkdir(parents=True, exist_ok=True)
        path = self.out_dir / filename
        path.write_bytes(data if isinstance(data, bytes) else data.encode("utf-8"))
        url = f"/scraped/{filename}" if self.out_dir.resolve() == self.scraped_dir.resolve() else None
        self.emit({"event": "file", "name": filename, "url": url, "path": str(path.resolve())})
        return f"Saved {filename}"

    def ask(self, question):
        """Pop a question up in the editor and wait for the kid's answer."""
        self._answer = None
        self._answered.clear()
        self.emit({"event": "ask", "question": question})
        while not self._answered.wait(0.2):
            self.check_stop()
        return self._answer

    def answer(self, text):
        self._answer = text
        self._answered.set()

    def fetch(self, url, limit=15 * 1024 * 1024):
        """Download bytes from the web (images, files). data: URLs are decoded directly."""
        if url.startswith("data:"):
            head, _, body = url.partition(",")
            return (base64.b64decode(body) if ";base64" in head else body.encode()), head[5:].split(";")[0]
        self.check_address(url)
        req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0 BlockBots"})
        try:
            with urllib.request.urlopen(req, timeout=20) as resp:
                data = resp.read(limit + 1)
                mime = resp.headers.get_content_type()
        except Exception as e:
            raise Oops(f"I couldn't download {short(url)} ({e}).")
        if len(data) > limit:
            raise Oops(f"That file is bigger than {limit // (1024 * 1024)} MB, so I didn't download it.")
        return data, mime

    # ---------- browser helpers ----------

    def open(self, name="chrome", hidden=False):
        if self.driver:
            return False
        if self.server_mode:
            self.log("Starting Chrome on the server. Watch it in the little browser window!")
            opts = webdriver.ChromeOptions()
            opts.unhandled_prompt_behavior = "ignore"  # popups stay open until a popup block handles them
            for arg in ("--headless=new", "--no-sandbox", "--disable-dev-shm-usage",
                        "--disable-gpu", "--window-size=1280,800", "--lang=en-US"):
                opts.add_argument(arg)
            if os.environ.get("CHROME_BIN"):
                opts.binary_location = os.environ["CHROME_BIN"]
            service = webdriver.ChromeService(executable_path=os.environ.get("CHROMEDRIVER") or None)
            self.driver = webdriver.Chrome(options=opts, service=service)
            self.live = True
            return True
        self.live = hidden
        self.log(f"Starting {name.title()}... (the very first time can take a minute)")
        if name == "firefox":
            opts = webdriver.FirefoxOptions()
            opts.unhandled_prompt_behavior = "ignore"
            if hidden:
                opts.add_argument("-headless")
            self.driver = webdriver.Firefox(options=opts)
        else:
            opts = webdriver.EdgeOptions() if name == "edge" else webdriver.ChromeOptions()
            opts.unhandled_prompt_behavior = "ignore"
            if hidden:
                opts.add_argument("--headless=new")
            if self.keep_open:
                opts.add_experimental_option("detach", True)
            opts.add_argument("--disable-search-engine-choice-screen")
            opts.add_experimental_option("prefs", {
                "credentials_enable_service": False,
                "profile.password_manager_enabled": False,
                "profile.password_manager_leak_detection": False,
            })
            self.driver = webdriver.Edge(options=opts) if name == "edge" else webdriver.Chrome(options=opts)
        return True

    @property
    def browser(self):
        if self.driver is None:
            self.log("No browser yet, so I'm opening Chrome for you.")
            self.open()
        return self.driver

    def wait(self, seconds=None):
        return WatchedWait(self, self.timeout if seconds is None else seconds)

    def locator(self, f, p=""):
        by = LOCATORS.get(f.get(p + "by", "css"), By.CSS_SELECTOR)
        value = str(f.get(p + "target", "")).strip()
        if not value:
            raise Oops("One of the element boxes is empty. Tell me what to look for!")
        return by, value

    def describe(self, f, p=""):
        by = f.get(p + "by", "css")
        value = str(f.get(p + "target", "")).strip()
        return f'"{value}"' if by == "css" else f'{by} "{value}"'

    def find(self, f, p="", cond=EC.presence_of_element_located, seconds=None):
        loc = self.locator(f, p)
        wait_for = self.timeout if seconds is None else seconds
        if self.server_mode:
            wait_for = min(wait_for, 60)
        try:
            el = self.wait(wait_for).until(cond(loc))
        except TimeoutException:
            raise Oops(f"I couldn't find {self.describe(f, p)} on the page (I waited {wait_for:g}s).")
        self.sparkle(el)
        return el

    def find_all(self, f, p=""):
        return self.browser.find_elements(*self.locator(f, p))

    def sparkle(self, el):
        """Flash a bright outline on the element so kids can see what the robot touches."""
        try:
            self.driver.execute_script(
                """const el = arguments[0], o = el.style.outline, s = el.style.boxShadow;
                el.style.outline = '4px solid #ff4fa3';
                el.style.boxShadow = '0 0 0 8px rgba(255,183,3,.6)';
                setTimeout(() => { el.style.outline = o; el.style.boxShadow = s; }, 450);""",
                el,
            )
        except Exception:
            pass

    def alert(self):
        try:
            return self.wait(5).until(EC.alert_is_present())
        except TimeoutException:
            raise Oops("I waited for a popup but none appeared.")


class WatchedWait(WebDriverWait):
    """WebDriverWait that keeps the live browser window updating, and lets Stop
    work, while it polls."""

    def __init__(self, run, seconds):
        super().__init__(run.browser, seconds)
        self._run = run

    def until(self, method, message=""):
        def watched(driver):
            self._run.check_stop()
            self._run.frame(force=False)
            return method(driver)
        return super().until(watched, message)


def compare(actual, mode, expected, what):
    actual, expected = str(actual).strip(), str(expected).strip()
    good = {
        "is": actual == expected,
        "contains": expected in actual,
        "not": actual != expected,
    }[mode]
    if not good:
        wanted = {"is": "be", "contains": "contain", "not": "NOT be"}[mode]
        raise Oops(f'{what} is "{actual}", but I expected it to {wanted} "{expected}".')
    return actual


# ================= Start =================

@block("start")
def start(run, f, kids):
    return "Ready, set, go!"


# ================= Browser =================

@block("open_browser")
def open_browser(run, f, kids):
    name = "chrome" if run.server_mode else f.get("browser", "chrome")
    if run.open(name, f.get("mode") == "hidden"):
        return f"Opened {name.title()}"
    return "The browser is already open"


@block("go_to")
def go_to(run, f, kids):
    if not str(f.get("url", "")).strip():
        raise Oops("Type a web address to go to.")
    url = web_address(f.get("url"))
    run.check_address(url)
    run.browser.get(url)
    return f"Went to {url}"


@block("back")
def back(run, f, kids):
    run.browser.back()
    return "Went back"


@block("forward")
def forward(run, f, kids):
    run.browser.forward()
    return "Went forward"


@block("refresh")
def refresh(run, f, kids):
    run.browser.refresh()
    return "Refreshed the page"


@block("maximize")
def maximize(run, f, kids):
    run.browser.maximize_window()
    return "Made the window big"


@block("window_size")
def window_size(run, f, kids):
    w, h = int(num(f.get("width"))), int(num(f.get("height")))
    run.browser.set_window_size(w, h)
    return f"Window is now {w}×{h}"


@block("new_tab")
def new_tab(run, f, kids):
    url = str(f.get("url", "")).strip()
    if url:
        run.check_address(web_address(url))
    run.browser.switch_to.new_window("tab")
    if url:
        run.browser.get(web_address(url))
    return f"Opened a new tab{' at ' + url if url else ''}"


@block("switch_tab")
def switch_tab(run, f, kids):
    handles = run.browser.window_handles
    i = int(num(f.get("index")))
    if not 1 <= i <= len(handles):
        raise Oops(f"There {'is' if len(handles) == 1 else 'are'} only {len(handles)} tab(s), so I can't go to tab #{i}.")
    run.browser.switch_to.window(handles[i - 1])
    return f"Switched to tab #{i}"


@block("close_tab")
def close_tab(run, f, kids):
    run.browser.close()
    handles = run.driver.window_handles
    if handles:
        run.driver.switch_to.window(handles[-1])
    return "Closed the tab"


@block("screenshot")
def screenshot(run, f, kids):
    safe = "".join(c for c in str(f.get("name", "picture")) if c.isalnum() or c in "-_") or "picture"
    file = f"{safe}-{datetime.now():%H%M%S}.png"
    if run.server_mode:  # nothing piles up on the server: the picture goes straight to the diary
        src = "data:image/png;base64," + run.browser.get_screenshot_as_base64()
        run.emit({"event": "screenshot", "name": safe, "src": src, "file": file})
    else:
        run.shots_dir.mkdir(parents=True, exist_ok=True)
        run.browser.save_screenshot(str(run.shots_dir / file))
        run.emit({"event": "screenshot", "name": safe, "url": f"/shots/{file}", "file": file})
    return f"Took a picture called {safe}"


@block("delete_cookies")
def delete_cookies(run, f, kids):
    run.browser.delete_all_cookies()
    return "Ate all the cookies 🍪"


@block("close_browser")
def close_browser(run, f, kids):
    if run.driver:
        run.driver.quit()
        run.driver = None
    return "Closed the browser"


# ================= Actions =================

@block("click")
def click(run, f, kids):
    run.find(f, cond=EC.element_to_be_clickable).click()
    return f"Clicked {run.describe(f)}"


@block("double_click")
def double_click(run, f, kids):
    ActionChains(run.browser).double_click(run.find(f, cond=EC.visibility_of_element_located)).perform()
    return f"Double-clicked {run.describe(f)}"


@block("right_click")
def right_click(run, f, kids):
    ActionChains(run.browser).context_click(run.find(f, cond=EC.visibility_of_element_located)).perform()
    return f"Right-clicked {run.describe(f)}"


@block("hover")
def hover(run, f, kids):
    ActionChains(run.browser).move_to_element(run.find(f, cond=EC.visibility_of_element_located)).perform()
    return f"Hovered over {run.describe(f)}"


@block("type_text")
def type_text(run, f, kids):
    text = str(f.get("text", ""))
    run.find(f, cond=EC.visibility_of_element_located).send_keys(text)
    return f'Typed "{text}" into {run.describe(f)}'


@block("clear_field")
def clear_field(run, f, kids):
    run.find(f, cond=EC.visibility_of_element_located).clear()
    return f"Cleared {run.describe(f)}"


@block("press_key")
def press_key(run, f, kids):
    key = f.get("key", "enter")
    if str(f.get("target", "")).strip():
        run.find(f, cond=EC.visibility_of_element_located).send_keys(KEYS[key])
        return f"Pressed {key.upper()} on {run.describe(f)}"
    ActionChains(run.browser).send_keys(KEYS[key]).perform()
    return f"Pressed {key.upper()}"


@block("select_text")
def select_text(run, f, kids):
    option = str(f.get("option", ""))
    try:
        Select(run.find(f)).select_by_visible_text(option)
    except NoSuchElementException:
        raise Oops(f'The dropdown {run.describe(f)} has no option called "{option}".')
    return f'Chose "{option}"'


@block("select_value")
def select_value(run, f, kids):
    value = str(f.get("value", ""))
    try:
        Select(run.find(f)).select_by_value(value)
    except NoSuchElementException:
        raise Oops(f'The dropdown {run.describe(f)} has no option with value "{value}".')
    return f'Chose value "{value}"'


@block("set_checkbox")
def set_checkbox(run, f, kids):
    el = run.find(f)
    want = f.get("state", "on") == "on"
    if el.is_selected() != want:
        el.click()
    return f"{'Ticked' if want else 'Unticked'} {run.describe(f)}"


@block("drag_drop")
def drag_drop(run, f, kids):
    src = run.find(f, cond=EC.visibility_of_element_located)
    dst = run.find(f, "to_", cond=EC.visibility_of_element_located)
    ActionChains(run.browser).click_and_hold(src).pause(0.2).move_to_element(dst).pause(0.2).release().perform()
    return f"Dragged {run.describe(f)} onto {run.describe(f, 'to_')}"


@block("scroll_to")
def scroll_to(run, f, kids):
    el = run.find(f)
    run.browser.execute_script("arguments[0].scrollIntoView({block: 'center', behavior: 'smooth'})", el)
    return f"Scrolled to {run.describe(f)}"


@block("scroll_by")
def scroll_by(run, f, kids):
    px = int(num(f.get("pixels")))
    run.browser.execute_script("window.scrollBy({top: arguments[0], behavior: 'smooth'})", px)
    return f"Scrolled {px} pixels"


@block("upload_file")
def upload_file(run, f, kids):
    if run.server_mode:
        raise Oops("Uploading files only works when the runner is on your own computer, not on the website.")
    path = Path(str(f.get("path", "")).strip().strip('"')).expanduser()
    if not path.is_file():
        raise Oops(f"I can't find the file {path} on this computer.")
    run.find(f).send_keys(str(path.resolve()))
    return f"Uploaded {path.name}"


@block("submit_form")
def submit_form(run, f, kids):
    run.find(f).submit()
    return f"Submitted {run.describe(f)}"


@block("run_js")
def run_js(run, f, kids):
    result = run.browser.execute_script(str(f.get("code", "")))
    return "Ran the JavaScript" if result is None else f"JavaScript answered: {result}"


# ================= Wait =================

@block("wait_seconds")
def wait_seconds(run, f, kids):
    s = num(f.get("seconds"))
    run.pause(s)
    return f"Waited {s:g} second(s)"


@block("wait_visible")
def wait_visible(run, f, kids):
    run.find(f, cond=EC.visibility_of_element_located, seconds=num(f.get("seconds")))
    return f"{run.describe(f)} showed up"


@block("wait_clickable")
def wait_clickable(run, f, kids):
    run.find(f, cond=EC.element_to_be_clickable, seconds=num(f.get("seconds")))
    return f"{run.describe(f)} is ready to click"


@block("wait_gone")
def wait_gone(run, f, kids):
    s = num(f.get("seconds"))
    try:
        run.wait(s).until(EC.invisibility_of_element_located(run.locator(f)))
    except TimeoutException:
        raise Oops(f"{run.describe(f)} was still there after {s:g}s.")
    return f"{run.describe(f)} disappeared"


@block("wait_title")
def wait_title(run, f, kids):
    text = str(f.get("text", ""))
    try:
        run.wait().until(EC.title_contains(text))
    except TimeoutException:
        raise Oops(f'The title never had "{text}" in it. It is "{run.browser.title}".')
    return f'Title has "{text}"'


@block("wait_text")
def wait_text(run, f, kids):
    text = str(f.get("text", ""))
    try:
        run.wait().until(EC.text_to_be_present_in_element(run.locator(f), text))
    except TimeoutException:
        raise Oops(f'{run.describe(f)} never said "{text}".')
    return f'{run.describe(f)} says "{text}"'


# ================= Check =================

@block("check_title")
def check_title(run, f, kids):
    compare(run.browser.title, f.get("mode", "is"), f.get("text", ""), "The title")
    return "Title looks right"


@block("check_url")
def check_url(run, f, kids):
    compare(run.browser.current_url, f.get("mode", "contains"), f.get("text", ""), "The web address")
    return "Web address looks right"


@block("check_text")
def check_text(run, f, kids):
    actual = compare(run.find(f).text, f.get("mode", "is"), f.get("text", ""), f"The text of {run.describe(f)}")
    return f'Text is "{actual}"'


@block("check_visible")
def check_visible(run, f, kids):
    if not run.find(f).is_displayed():
        raise Oops(f"{run.describe(f)} is on the page but hidden.")
    return f"{run.describe(f)} is visible"


@block("check_hidden")
def check_hidden(run, f, kids):
    if any(el.is_displayed() for el in run.find_all(f)):
        raise Oops(f"{run.describe(f)} is visible, but I expected it to be hidden.")
    return f"{run.describe(f)} is hidden"


@block("check_enabled")
def check_enabled(run, f, kids):
    if not run.find(f).is_enabled():
        raise Oops(f"{run.describe(f)} is disabled.")
    return f"{run.describe(f)} is enabled"


@block("check_selected")
def check_selected(run, f, kids):
    if not run.find(f).is_selected():
        raise Oops(f"{run.describe(f)} is not ticked.")
    return f"{run.describe(f)} is ticked"


@block("check_attribute")
def check_attribute(run, f, kids):
    attr = str(f.get("attr", ""))
    compare(run.find(f).get_attribute(attr) or "", "is", f.get("value", ""), f'The "{attr}" of {run.describe(f)}')
    return f'"{attr}" looks right'


@block("check_count")
def check_count(run, f, kids):
    want, got = int(num(f.get("count"))), len(run.find_all(f))
    if got != want:
        raise Oops(f"I found {got} of {run.describe(f)}, but expected {want}.")
    return f"Found exactly {want}"


@block("check_value")
def check_value(run, f, kids):
    compare(run.find(f).get_attribute("value") or "", "is", f.get("value", ""), f"The value in {run.describe(f)}")
    return "Value looks right"


# ================= Scrape =================

def short(text, n=60):
    text = " ".join(str(text).split())
    return text if len(text) <= n else text[: n - 1] + "…"


def text_of(el):
    """Visible text, or the typed value for text boxes."""
    return (el.text or el.get_attribute("value") or "").strip()


@block("scrape_text")
def scrape_text(run, f, kids):
    text = text_of(run.find(f))
    name = run.remember(f.get("name"), text, card=True)
    return f'Grabbed "{short(text)}" as {name}'


@block("scrape_all")
def scrape_all(run, f, kids):
    run.find(f)  # wait until at least one is there
    texts = [t for t in (text_of(el) for el in run.find_all(f)) if t]
    name = run.remember(f.get("name"), texts, card=True)
    return f"Grabbed {len(texts)} text(s) as {name}"


@block("scrape_attribute")
def scrape_attribute(run, f, kids):
    attr = str(f.get("attr", "")).strip()
    if not attr:
        raise Oops("Type which attribute to grab, like href or src.")
    els = run.find_all(f) if f.get("which") == "all" else [run.find(f)]
    if not els:
        run.find(f)  # raises the friendly "couldn't find" message after waiting
    values = [el.get_attribute(attr) or "" for el in els]
    value = values if f.get("which") == "all" else values[0]
    name = run.remember(f.get("name"), value, card=True)
    return f'Grabbed {len(values)} "{attr}" value(s) as {name}' if isinstance(value, list) else f'Grabbed {attr} "{short(value)}" as {name}'


@block("scrape_table")
def scrape_table(run, f, kids):
    rows = run.browser.execute_script(
        "return [...arguments[0].querySelectorAll('tr')].map(r => [...r.cells].map(c => c.innerText.trim()))",
        run.find(f),
    )
    rows = [r for r in rows if any(r)]
    if not rows:
        raise Oops(f"{run.describe(f)} has no table rows to grab.")
    has_header = run.browser.execute_script("return !!arguments[0].querySelector('th')", run.find(f))
    columns = rows.pop(0) if has_header else [f"Column {i + 1}" for i in range(max(map(len, rows)))]
    name = run.remember(f.get("name"), {"columns": columns, "rows": rows}, card=True)
    return f"Grabbed a table with {len(rows)} row(s) as {name}"


@block("scrape_links")
def scrape_links(run, f, kids):
    rows = run.browser.execute_script(
        "return [...document.links].filter(a => a.href).map(a => [a.innerText.trim() || a.title || '(no text)', a.href])"
    )
    name = run.remember(f.get("name"), {"columns": ["text", "link"], "rows": rows}, card=True)
    return f"Grabbed {len(rows)} link(s) as {name}"


@block("scrape_page")
def scrape_page(run, f, kids):
    what = f.get("what", "title")
    value = {
        "title": lambda: run.browser.title,
        "url": lambda: run.browser.current_url,
        "text": lambda: run.browser.find_element(By.TAG_NAME, "body").text,
    }[what]()
    name = run.remember(f.get("name"), value, card=True)
    return f'Grabbed the page {"words" if what == "text" else what} as {name}'


@block("scrape_count")
def scrape_count(run, f, kids):
    count = len(run.find_all(f))
    name = run.remember(f.get("name"), count, card=True)
    return f"Counted {count} of {run.describe(f)} as {name}"


@block("check_scraped")
def check_scraped(run, f, kids):
    name, value = run.recall(f.get("name"))
    mode, expected = f.get("mode", "is"), f.get("text", "")
    if mode in ("more", "less"):
        want = number_in(expected)
        word = "more" if mode == "more" else "less"
        ok = (lambda n: n > want) if mode == "more" else (lambda n: n < want)
        if isinstance(value, list):  # every item in a list has to pass
            bad = [v for v in value if not ok(number_in(v))]
            if bad:
                raise Oops(f'Not every item in {name} is {word} than {want:g}: "{bad[0]}" isn\'t.')
            return f"All {len(value)} items in {name} are {word} than {want:g}"
        got = value if isinstance(value, (int, float)) else number_in(as_text(value))
        if ok(got):
            return f"{name} is {got:g}, {word} than {want:g}"
        raise Oops(f"{name} is {got:g}, but I expected it to be {word} than {want:g}.")
    compare(as_text(value), mode, expected, name)
    return f"{name} looks right"


# ================= Files (inputs & outputs) =================

def stamp(name, ext):
    """my-file -> my-file-153012.csv, so saving twice never overwrites."""
    return f"{safe_filename(name or 'data')}-{datetime.now():%H%M%S}.{ext}"


XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@block("save_scraped")
def save_scraped(run, f, kids):
    name, value = run.recall(f.get("name"))
    fmt = f.get("format", "csv")
    filename = stamp(f.get("file") or name, fmt)
    if fmt == "json":
        data, mime = json.dumps(value, indent=2, ensure_ascii=False), "application/json"
    elif fmt == "xlsx":
        data, mime = to_xlsx({name: rows_of(name, value)}), XLSX
    elif fmt == "txt":
        items = value["rows"] if isinstance(value, dict) else value if isinstance(value, list) else [value]
        data, mime = "\n".join(as_text(v) for v in items), "text/plain"
    else:
        data, mime = to_csv(rows_of(name, value)), "text/csv"
    return run.output(filename, data, mime)


@block("save_all")
def save_all(run, f, kids):
    data = {k: v for k, v in run.vars.items() if k not in ("item", "number")}
    if not data:
        raise Oops("There are no variables with values to save yet.")
    if f.get("format") == "json":
        return run.output(stamp(f.get("file"), "json"), json.dumps(data, indent=2, ensure_ascii=False), "application/json")
    simple = [["variable", "value"]] + [[k, v] for k, v in data.items() if not isinstance(v, (list, dict))]
    sheets = {"Variables": simple, **{k: rows_of(k, v) for k, v in data.items() if isinstance(v, (list, dict))}}
    return run.output(stamp(f.get("file"), "xlsx"), to_xlsx(sheets), XLSX)


@block("output_folder")
def output_folder(run, f, kids):
    if run.server_mode:
        return "On the website files are downloaded to your computer instead"
    raw = str(f.get("path", "")).strip().strip('"')
    if not raw:
        raise Oops("Type a folder, like C:\\Users\\me\\Documents\\robot.")
    folder = Path(raw).expanduser()
    folder.mkdir(parents=True, exist_ok=True)
    run.out_dir = folder
    return f"Files will be saved in {folder}"


@block("load_file")
def load_file(run, f, kids):
    picked = f.get("file") or {}
    if not isinstance(picked, dict) or not picked.get("data"):
        raise Oops("Click 📁 in the block to choose a file first.")
    value = read_table(picked["name"], base64.b64decode(picked["data"]))
    name = run.remember(f.get("name"), value, card=True)
    size = f"{len(value['rows'])} rows" if isinstance(value, dict) else f"{len(value)} items" if isinstance(value, list) else "data"
    return f"Loaded {picked['name']} ({size}) into {name}"


def extension(url, mime):
    last = urlparse(url).path.rsplit("/", 1)[-1]
    if "." in last:
        return last.rsplit(".", 1)[-1][:5].lower()
    return {"jpeg": "jpg", "svg+xml": "svg", "plain": "txt"}.get(mime.split("/")[-1], mime.split("/")[-1][:5] or "bin")


@block("download_image")
def download_image(run, f, kids):
    el = run.find(f)
    src = run.browser.execute_script("return arguments[0].currentSrc || arguments[0].src || ''", el)
    if src:
        data, mime = run.fetch(src)
        ext = extension(src, mime) if not src.startswith("data:") else extension("", mime)
    else:  # not an <img>: take a picture of it instead
        data, mime, ext = el.screenshot_as_png, "image/png", "png"
    return run.output(stamp(f.get("file"), ext), data, mime)


@block("download_link")
def download_link(run, f, kids):
    el = run.find(f)
    href = el.get_attribute("href") or el.get_attribute("src")
    if not href:
        raise Oops(f"{run.describe(f)} doesn't link to a file.")
    data, mime = run.fetch(href)
    return run.output(stamp(f.get("file"), extension(href, mime)), data, mime)


@block("download_url")
def download_url(run, f, kids):
    url = web_address(f.get("url", ""))
    data, mime = run.fetch(url)
    return run.output(stamp(f.get("file"), extension(url, mime)), data, mime)


@block("element_picture")
def element_picture(run, f, kids):
    png = run.find(f, cond=EC.visibility_of_element_located).screenshot_as_png
    file = stamp(f.get("file"), "png")
    run.emit({"event": "screenshot", "name": f.get("file") or "element",
              "src": "data:image/png;base64," + base64.b64encode(png).decode(), "file": file})
    return run.output(file, png, "image/png")


@block("full_page_picture")
def full_page_picture(run, f, kids):
    d = run.browser
    if hasattr(d, "get_full_page_screenshot_as_png"):  # Firefox
        png = d.get_full_page_screenshot_as_png()
    else:
        shot = d.execute_cdp_cmd("Page.captureScreenshot", {"format": "png", "captureBeyondViewport": True})
        png = base64.b64decode(shot["data"])
    return run.output(stamp(f.get("file"), "png"), png, "image/png")


@block("save_pdf")
def save_pdf(run, f, kids):
    try:
        pdf = base64.b64decode(run.browser.print_page())
    except WebDriverException:
        raise Oops('Saving a PDF only works when the browser is hidden. Choose "🙈 hidden" in the open browser block.')
    return run.output(stamp(f.get("file"), "pdf"), pdf, "application/pdf")


@block("save_html")
def save_html(run, f, kids):
    return run.output(stamp(f.get("file"), "html"), run.browser.page_source, "text/html")


# ================= Popups & frames =================

@block("switch_frame")
def switch_frame(run, f, kids):
    run.browser.switch_to.frame(run.find(f))
    return f"Went inside frame {run.describe(f)}"


@block("switch_main")
def switch_main(run, f, kids):
    run.browser.switch_to.default_content()
    return "Back on the main page"


@block("accept_alert")
def accept_alert(run, f, kids):
    alert = run.alert()
    text = alert.text
    alert.accept()
    return f'Clicked OK on "{text}"'


@block("dismiss_alert")
def dismiss_alert(run, f, kids):
    alert = run.alert()
    text = alert.text
    alert.dismiss()
    return f'Clicked Cancel on "{text}"'


@block("type_alert")
def type_alert(run, f, kids):
    alert = run.alert()
    alert.send_keys(str(f.get("text", "")))
    alert.accept()
    return f'Typed "{f.get("text", "")}" into the popup'


@block("check_alert")
def check_alert(run, f, kids):
    compare(run.alert().text, "contains", f.get("text", ""), "The popup")
    return "Popup says the right thing"


# ================= Control =================

@block("repeat")
def repeat(run, f, kids):
    times = int(num(f.get("times")))
    for _ in range(times):
        run.run_blocks(kids)
    return f"Repeated {times} time(s)"


@block("if_exists")
def if_exists(run, f, kids):
    if any(el.is_displayed() for el in run.find_all(f)):
        run.run_blocks(kids)
        return f"Found {run.describe(f)}, so I did the inside blocks"
    return f"No {run.describe(f)}, so I skipped the inside blocks"


@block("say")
def say(run, f, kids):
    return f"💬 {f.get('message', '')}"


@block("for_each")
def for_each(run, f, kids):
    """Run the inside blocks once per item. For a table, every column also becomes a
    variable for the current row, e.g. {Email} - handy for data-driven tests."""
    name, value = run.recall(f.get("name"))
    columns = [var_name(c) for c in value["columns"]] if isinstance(value, dict) else []
    items = value["rows"] if isinstance(value, dict) else value if isinstance(value, list) else [value]
    for i, item in enumerate(items, 1):
        run.remember("number", i)
        run.remember("item", item)
        for col, cell in zip(columns, item if columns else []):
            if col:
                run.remember(col, cell)
        run.run_blocks(kids)
    return f"Went through all {len(items)} item(s) in {name}"


def comparison(a, op, b):
    try:
        a, b = as_number(a), as_number(b)  # compare as numbers when both are numbers
    except Oops:
        a, b = as_text(a).strip(), as_text(b).strip()
        if op in (">", "<"):
            raise Oops(f'To use "{op}", both sides need numbers (got "{a}" and "{b}").')
    return {"=": a == b, "!=": a != b, ">": a > b, "<": a < b, "contains": as_text(b) in as_text(a)}[op]


@block("if_compare")
def if_compare(run, f, kids):
    a, op, b = f.get("a", ""), f.get("op", "="), f.get("b", "")
    said = f'"{short(as_text(a), 30)}" {op} "{short(as_text(b), 30)}"'
    if comparison(a, op, b):
        run.run_blocks(kids)
        return f"{said} is true, so I did the inside blocks"
    return f"{said} is false, so I skipped the inside blocks"


@block("repeat_until")
def repeat_until(run, f, kids):
    raw = run.current_raw  # re-read the fields each time so variables changed inside count
    for i in range(1, 1001):
        a, b = run.fill(raw.get("a", "")), run.fill(raw.get("b", ""))
        if comparison(a, raw.get("op", "="), b):
            return f"Stopped repeating after {i - 1} time(s)"
        run.run_blocks(kids)
    raise Oops("I repeated 1000 times and it still wasn't true, so I stopped.")


# ================= Variables =================

@block("set_var")
def set_var(run, f, kids):
    value = f.get("value", "")
    if re.fullmatch(r"\s*-?\d+(\.\d+)?\s*", str(value)):
        value = as_number(value)
    name = run.remember(f.get("name"), value)
    return f"{name} = {short(as_text(value))}"


@block("change_var")
def change_var(run, f, kids):
    name = var_name(f.get("name"))
    new = as_number(as_number(run.vars.get(name, 0)) + as_number(f.get("by", 1)))
    run.remember(name, new)
    return f"{name} is now {new}"


@block("math_var")
def math_var(run, f, kids):
    a, b, op = as_number(f.get("a")), as_number(f.get("b")), f.get("op", "+")
    if op in ("/", "%") and b == 0:
        raise Oops("You can't divide by zero!")
    result = as_number(round({"+": a + b, "-": a - b, "*": a * b, "/": a / b, "%": a % b}[op], 6))
    name = run.remember(f.get("name"), result)
    return f"{name} = {a} {op} {b} = {result}"


@block("join_var")
def join_var(run, f, kids):
    value = as_text(f.get("a", "")) + as_text(f.get("b", ""))
    name = run.remember(f.get("name"), value)
    return f'{name} = "{short(value)}"'


@block("random_var")
def random_var(run, f, kids):
    lo, hi = sorted((int(as_number(f.get("from", 1))), int(as_number(f.get("to", 10)))))
    value = random.randint(lo, hi)
    name = run.remember(f.get("name"), value)
    return f"{name} = {value} (random)"


@block("date_var")
def date_var(run, f, kids):
    formats = {"date": "%Y-%m-%d", "time": "%H:%M:%S", "datetime": "%Y-%m-%d %H:%M:%S", "stamp": "%Y%m%d-%H%M%S"}
    value = datetime.now().strftime(formats.get(f.get("what"), formats["date"]))
    name = run.remember(f.get("name"), value)
    return f"{name} = {value}"


@block("list_add")
def list_add(run, f, kids):
    name = var_name(f.get("name"))
    items = list(run.vars[name]) if isinstance(run.vars.get(name), list) else []
    items.append(f.get("value", ""))
    run.remember(name, items)
    return f"{name} now has {len(items)} item(s)"


@block("list_clear")
def list_clear(run, f, kids):
    name = run.remember(f.get("name"), [])
    return f"{name} is now an empty list"


@block("ask_var")
def ask_var(run, f, kids):
    answer = run.ask(str(f.get("question", "")))
    name = run.remember(f.get("name"), answer)
    return f'You answered "{short(answer)}", saved in {name}'
